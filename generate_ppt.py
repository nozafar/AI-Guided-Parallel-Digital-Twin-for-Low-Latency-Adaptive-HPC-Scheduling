from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    
    # Slide Dimensions (16:9 widescreen)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Theme Colors (Dark AI / HPC Theme)
    BG_COLOR = RGBColor(20, 22, 28)       # Dark Slate
    TITLE_COLOR = RGBColor(0, 229, 255)   # Neon Cyan
    BODY_COLOR = RGBColor(230, 235, 240)  # Off-White
    ACCENT_COLOR = RGBColor(255, 0, 102)  # Neon Pink
    
    # Define slide layouts (blank layouts so we can custom design them)
    blank_slide_layout = prs.slide_layouts[6] 

    def apply_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_COLOR
        
        # Add decorative top accent line
        accent = slide.shapes.add_shape(
            1, # MSO_SHAPE.RECTANGLE
            Inches(0), Inches(0), prs.slide_width, Inches(0.1)
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = ACCENT_COLOR
        accent.line.fill.background()

    def add_title_slide(title_text, subtitle_text):
        slide = prs.slides.add_slide(blank_slide_layout)
        apply_background(slide)
        
        # Title
        txBox = slide.shapes.add_textbox(Inches(1), Inches(2.2), Inches(11.33), Inches(1.5))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.add_paragraph()
        p.text = title_text
        p.font.bold = True
        p.font.size = Pt(64)
        p.font.color.rgb = TITLE_COLOR
        p.font.name = 'Arial'
        p.alignment = PP_ALIGN.CENTER
        
        # Subtitle
        subBox = slide.shapes.add_textbox(Inches(1), Inches(4.2), Inches(11.33), Inches(1.5))
        tf_sub = subBox.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.add_paragraph()
        p_sub.text = subtitle_text
        p_sub.font.size = Pt(28)
        p_sub.font.color.rgb = BODY_COLOR
        p_sub.font.name = 'Arial'
        p_sub.alignment = PP_ALIGN.CENTER
        
        return slide

    def add_slide(title_text, content):
        slide = prs.slides.add_slide(blank_slide_layout)
        apply_background(slide)
        
        # Title
        txBox = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1))
        tf = txBox.text_frame
        p = tf.add_paragraph()
        p.text = title_text
        p.font.bold = True
        p.font.size = Pt(40)
        p.font.color.rgb = TITLE_COLOR
        p.font.name = 'Arial'
        
        # Body - Adjusted dimensions to prevent overflow
        bodyBox = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.2))
        tf_body = bodyBox.text_frame
        tf_body.word_wrap = True
        
        for idx, point in enumerate(content):
            p = tf_body.add_paragraph()
            p.text = "•  " + point
            p.font.size = Pt(22) # Reduced from 26 to prevent text overflow
            p.font.color.rgb = BODY_COLOR
            p.font.name = 'Arial'
            p.space_after = Pt(12) # Reduced spacing to fit more text safely
            
        return slide

    # Slide 1: Title
    add_title_slide(
        "AI-HPC Digital Twin",
        "Intelligent Job Scheduling using Reinforcement Learning\n\nSystem Architecture & Research Overview"
    )

    # Slide 2: Abstract
    add_slide("Abstract", [
        "High-Performance Computing (HPC) clusters suffer from inefficient resource allocation.",
        "Traditional schedulers (like FCFS or SJF) are rigid and fail to adapt to dynamic workloads.",
        "This project introduces an AI-driven Digital Twin for HPC scheduling.",
        "We simulate cluster environments and use Reinforcement Learning (PPO) to learn optimal policies.",
        "Goal: Minimize wait times, maximize resource utilization, and prevent deadlocks autonomously."
    ])

    # Slide 3: Introduction
    add_slide("Introduction: The HPC Environment", [
        "HPC clusters consist of massive compute nodes with CPUs, GPUs, and Memory.",
        "Thousands of jobs arrive continuously with varying resource demands.",
        "The Scheduler is the 'brain' of the cluster, deciding WHICH job runs WHEN.",
        "Poor scheduling leads to idle GPUs, blocked pipelines, and wasted millions in energy costs."
    ])

    # Slide 4: What is the Core Question?
    add_slide("The Core Problem Statement", [
        "QUESTION: How can we dynamically adapt to unpredictable HPC workloads better than static algorithms?",
        "Sub-problems:",
        "How to avoid cluster fragmentation (bin-packing problem)?",
        "How to prioritize jobs without starving long-running tasks?",
        "How to safely test new algorithms without risking a $10M real-world supercomputer?"
    ])

    # Slide 5: Why This Project?
    add_slide("Why This Project?", [
        "Real-world testing is impossible: You cannot pause a supercomputer to test a new algorithm.",
        "Digital Twin solves this: We built a 1:1 simulation environment (SimPy) mimicking an HPC.",
        "AI provides adaptability: RL algorithms can discover 'cheat codes' in scheduling that humans miss.",
        "Bridging domains: Combining Systems Engineering (HPC) with advanced Deep Learning (RL)."
    ])

    # Slide 6: Traditional Schedulers (The Baseline)
    add_slide("The Flaws of Traditional Methods", [
        "First-Come, First-Serve (FCFS): Simple, but a massive job blocks the entire cluster.",
        "Shortest Job First (SJF): Good for overall throughput, but massive jobs get starved forever.",
        "Priority Queues: Often abused by users, hard to balance fairness.",
        "None of these learn from past mistakes or adapt to current load."
    ])

    # Slide 7: What Our Model Does
    add_slide("What Our System Does", [
        "Generates or ingests real workload traces (e.g., Alibaba GPU traces).",
        "Pushes jobs into the Digital Twin (SimPy simulator).",
        "The RL Agent (PPO) observes the cluster state (queue length, available CPUs/GPUs).",
        "The Agent makes a decision: Schedule, Delay, or Skip.",
        "The system rewards or penalizes the agent based on wait times and makespan."
    ])

    # Slide 8: The Digital Twin Architecture
    add_slide("The Digital Twin Architecture", [
        "The Cluster: Configurable nodes (e.g., 64 nodes, 32 CPUs, 4 GPUs each).",
        "The Simulator (SimPy): Discrete event simulation engine advancing time only when necessary.",
        "The Job Pipeline: Simulates arrival times, execution, and resource locking.",
        "Metrics Collector: Tracks CPU/GPU utilization and queue health in real-time."
    ])

    # Slide 9: Reinforcement Learning Setup
    add_slide("Methodology: RL Environment", [
        "Observation Space (24 dimensions):",
        " - Cluster available resources (CPU, GPU, RAM).",
        " - Queue length and exact resource demands of the next 5 pending jobs.",
        "Action Space (Discrete):",
        " - Action 0: Schedule next job.",
        " - Action 1: Delay (Wait for running jobs to finish).",
        " - Action 2: Skip first job, schedule the second (solves head-of-line blocking)."
    ])

    # Slide 10: The Reward Function (The Secret Sauce)
    add_slide("Methodology: Reward Engineering", [
        "A poorly designed reward leads to deadlocks (agent decides to delay forever).",
        "Our Reward Structure:",
        " - Penalty: -0.1 per job waiting in the queue per second.",
        " - Penalty: -5 for attempting an illegal schedule (job doesn't fit).",
        " - Massive Penalty: -100,000 for deadlocking the cluster.",
        " - Bonus: +100 for successfully finishing a job."
    ])

    # Slide 11: Training the Agent
    add_slide("Training the PPO Agent", [
        "Framework: Stable-Baselines3 (Proximal Policy Optimization).",
        "Hardware: CUDA accelerated training.",
        "Episodes: Jobs are chunked into episodes of 100-1000 jobs.",
        "Learning process: The agent explores randomly at first, hits deadlocks, learns the penalties, and eventually figures out optimal bin-packing."
    ])

    # Slide 12: Experimental Results
    add_slide("Evaluation and Results", [
        "We evaluate the trained RL model against FCFS, SJF, and Priority schedulers.",
        "Metrics Tracked:",
        " - Average Waiting Time (How long a job sits in queue).",
        " - Makespan (Total time to finish all jobs).",
        " - Resource Utilization (CPU/GPU uptime).",
        "The simulation pipeline generates automated analytics and interactive dashboards."
    ])

    # Slide 13: Overcoming Challenges
    add_slide("Challenges & Solutions", [
        "Challenge: Time Advancement Logic.",
        "Solution: Adjusted the simulation step so the RL agent can schedule multiple jobs simultaneously.",
        "Challenge: Deadlocks during evaluation (Agent predicts 'Delay' infinitely).",
        "Solution: Implemented a robust fallback mechanism. If the cluster is completely idle, it forces a schedule, guaranteeing progress."
    ])

    # Slide 14: Future Work
    add_slide("Future Work", [
        "Scale the simulation to millions of jobs using the full 350GB Alibaba GPU trace.",
        "Action Space Expansion: Allow the agent to dynamically select exactly which nodes to allocate to (topology-aware scheduling).",
        "Multi-Agent RL: Have separate agents for CPU-heavy vs GPU-heavy scheduling.",
        "Predictive Scheduling: Integrate XGBoost runtime predictions directly into the RL state."
    ])

    # Slide 15: Conclusion
    add_slide("Conclusion", [
        "HPC scheduling is too complex for static heuristic algorithms.",
        "Digital Twins allow zero-risk, high-speed iteration of cluster management software.",
        "Reinforcement Learning, when constrained with robust fallback mechanics and shaped rewards, provides a dynamic, self-optimizing scheduler.",
        "This project lays the groundwork for next-generation AI-managed supercomputers."
    ])

    prs.save('AI_HPC_Digital_Twin_Presentation_Designed_v2.pptx')
    print("Beautiful Presentation saved successfully as AI_HPC_Digital_Twin_Presentation_Designed_v2.pptx")

if __name__ == '__main__':
    create_presentation()
